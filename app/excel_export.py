import re

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

_FORBIDDEN_IN_SHEET_NAME = re.compile(r"[\[\]:*?/\\]")


def _sheet_title(name: str, used: set) -> str:
    base = _FORBIDDEN_IN_SHEET_NAME.sub(" ", name).strip()[:31] or "Thème"
    title = base
    n = 2
    while title.lower() in used:
        suffix = f" ({n})"
        title = base[: 31 - len(suffix)] + suffix
        n += 1
    used.add(title.lower())
    return title


def export_questions_to_excel(db, path: str) -> tuple:
    """Write every theme as one sheet (columns Question / Réponse).

    The layout is the same one the import dialogs (PC and mobile) read.
    Returns (number of themes, number of questions).
    """
    wb = Workbook()
    wb.remove(wb.active)
    used: set = set()
    total_questions = 0
    themes = db.list_themes()

    for theme in themes:
        ws = wb.create_sheet(_sheet_title(theme.name, used))
        ws.append(["Question", "Réponse"])
        for q in db.list_questions(theme.id):
            ws.append([q.question, q.answer])
            total_questions += 1
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="2F5597")
        ws.column_dimensions["A"].width = 80
        ws.column_dimensions["B"].width = 38
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.freeze_panes = "A2"

    if not themes:
        ws = wb.create_sheet("Questions")
        ws.append(["Question", "Réponse"])

    wb.save(path)
    return len(themes), total_questions
