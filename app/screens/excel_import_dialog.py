from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from PySide6.QtWidgets import (
    QComboBox,
    QDialog,
    QFileDialog,
    QFormLayout,
    QHBoxLayout,
    QLabel,
    QLineEdit,
    QMessageBox,
    QPushButton,
    QTableWidget,
    QTableWidgetItem,
    QVBoxLayout,
)

NEW_THEME_SENTINEL = "__new_theme__"


class ExcelImportDialog(QDialog):
    def __init__(self, db, parent=None):
        super().__init__(parent)
        self.db = db
        self.workbook = None
        self.file_path = None
        self.setWindowTitle("Importer des questions depuis un fichier Excel")
        self.resize(680, 560)

        root = QVBoxLayout(self)

        file_row = QHBoxLayout()
        self.lbl_file = QLabel("Aucun fichier sélectionné")
        btn_choose = QPushButton("Choisir un fichier Excel…")
        btn_choose.clicked.connect(self._choose_file)
        file_row.addWidget(btn_choose)
        file_row.addWidget(self.lbl_file, 1)
        root.addLayout(file_row)

        form = QFormLayout()
        self.combo_sheet = QComboBox()
        self.combo_sheet.currentIndexChanged.connect(self._on_sheet_changed)
        form.addRow("Feuille :", self.combo_sheet)

        self.combo_question_col = QComboBox()
        self.combo_question_col.currentIndexChanged.connect(self._update_preview)
        form.addRow("Colonne question :", self.combo_question_col)

        self.combo_answer_col = QComboBox()
        self.combo_answer_col.currentIndexChanged.connect(self._update_preview)
        form.addRow("Colonne réponse :", self.combo_answer_col)

        self.combo_theme = QComboBox()
        self.combo_theme.currentIndexChanged.connect(self._on_theme_choice_changed)
        form.addRow("Thème cible :", self.combo_theme)

        self.edit_new_theme = QLineEdit()
        self.edit_new_theme.setPlaceholderText("Nom du nouveau thème")
        self.edit_new_theme.setVisible(False)
        form.addRow("", self.edit_new_theme)

        root.addLayout(form)

        root.addWidget(QLabel("Aperçu (10 premières lignes) :"))
        self.preview = QTableWidget(0, 2)
        self.preview.setHorizontalHeaderLabels(["Question", "Réponse"])
        root.addWidget(self.preview, 1)

        buttons_row = QHBoxLayout()
        btn_import = QPushButton("Importer")
        btn_import.clicked.connect(self._do_import)
        btn_close = QPushButton("Fermer")
        btn_close.clicked.connect(self.accept)
        buttons_row.addStretch()
        buttons_row.addWidget(btn_import)
        buttons_row.addWidget(btn_close)
        root.addLayout(buttons_row)

        self._reload_theme_combo()
        self._set_controls_enabled(False)

    def _set_controls_enabled(self, enabled: bool):
        for w in (self.combo_sheet, self.combo_question_col, self.combo_answer_col, self.combo_theme):
            w.setEnabled(enabled)

    def _reload_theme_combo(self):
        self.combo_theme.blockSignals(True)
        self.combo_theme.clear()
        for theme in self.db.list_themes():
            self.combo_theme.addItem(theme.name, theme.id)
        self.combo_theme.addItem("+ Nouveau thème…", NEW_THEME_SENTINEL)
        self.combo_theme.blockSignals(False)

    def _on_theme_choice_changed(self):
        is_new = self.combo_theme.currentData() == NEW_THEME_SENTINEL
        self.edit_new_theme.setVisible(is_new)

    def _choose_file(self):
        path, _ = QFileDialog.getOpenFileName(self, "Choisir un fichier Excel", "", "Fichiers Excel (*.xlsx)")
        if not path:
            return
        try:
            self.workbook = load_workbook(path, read_only=True, data_only=True)
        except Exception as exc:
            QMessageBox.critical(self, "Erreur", f"Impossible de lire ce fichier :\n{exc}")
            return
        self.file_path = path
        self.lbl_file.setText(Path(path).name)
        self.combo_sheet.blockSignals(True)
        self.combo_sheet.clear()
        self.combo_sheet.addItems(self.workbook.sheetnames)
        self.combo_sheet.blockSignals(False)
        self._set_controls_enabled(True)
        self._on_sheet_changed()

    def _current_sheet(self):
        if not self.workbook:
            return None
        name = self.combo_sheet.currentText()
        if not name:
            return None
        return self.workbook[name]

    def _on_sheet_changed(self):
        sheet = self._current_sheet()
        if sheet is None:
            return
        header_row = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
        self.combo_question_col.blockSignals(True)
        self.combo_answer_col.blockSignals(True)
        self.combo_question_col.clear()
        self.combo_answer_col.clear()
        for i, value in enumerate(header_row, start=1):
            letter = get_column_letter(i)
            label = str(value).strip() if value not in (None, "") else "(sans en-tête)"
            text = f"{letter} — {label}"
            self.combo_question_col.addItem(text, i)
            self.combo_answer_col.addItem(text, i)
        if self.combo_answer_col.count() > 1:
            self.combo_answer_col.setCurrentIndex(1)
        self.combo_question_col.blockSignals(False)
        self.combo_answer_col.blockSignals(False)
        self._update_preview()

    def _iter_data_rows(self, sheet, q_col, a_col):
        for row in sheet.iter_rows(min_row=2, values_only=True):
            if q_col - 1 >= len(row) or a_col - 1 >= len(row):
                continue
            question = row[q_col - 1]
            answer = row[a_col - 1]
            question = str(question).strip() if question is not None else ""
            answer = str(answer).strip() if answer is not None else ""
            if not question or not answer:
                continue
            yield question, answer

    def _update_preview(self):
        self.preview.setRowCount(0)
        sheet = self._current_sheet()
        q_col = self.combo_question_col.currentData()
        a_col = self.combo_answer_col.currentData()
        if sheet is None or q_col is None or a_col is None:
            return
        rows = []
        for question, answer in self._iter_data_rows(sheet, q_col, a_col):
            rows.append((question, answer))
            if len(rows) >= 10:
                break
        self.preview.setRowCount(len(rows))
        for r, (q, a) in enumerate(rows):
            self.preview.setItem(r, 0, QTableWidgetItem(q))
            self.preview.setItem(r, 1, QTableWidgetItem(a))

    def _do_import(self):
        sheet = self._current_sheet()
        q_col = self.combo_question_col.currentData()
        a_col = self.combo_answer_col.currentData()
        if sheet is None or q_col is None or a_col is None:
            QMessageBox.warning(self, "Import", "Choisissez d'abord un fichier, une feuille et des colonnes.")
            return
        if q_col == a_col:
            QMessageBox.warning(self, "Import", "La colonne question et la colonne réponse doivent être différentes.")
            return

        theme_data = self.combo_theme.currentData()
        if theme_data == NEW_THEME_SENTINEL:
            name = self.edit_new_theme.text().strip()
            if not name:
                QMessageBox.warning(self, "Import", "Indiquez un nom pour le nouveau thème.")
                return
            theme = self.db.get_or_create_theme(name)
            theme_id = theme.id
        else:
            theme_id = theme_data

        pairs = list(self._iter_data_rows(sheet, q_col, a_col))
        count = self.db.bulk_add_questions(theme_id, pairs)
        self._reload_theme_combo()
        QMessageBox.information(self, "Import terminé", f"{count} question(s) importée(s).")
